"""実運用を想定したテストデータを作成する。

  python tests/make_testdata.py <出力フォルダ>

「受注管理システム」の開発フォルダを想定し、設計書 (Word/Excel/PowerPoint/PDF の新旧形式)、
ソース・設定ファイル (各種文字コード)、異常系ファイル (パスワード付き・破損・画像だけのPDF・
バイナリ・Office一時ファイル) を作る。Word・Excel には図形 (テキストボックス・グループ図形・
旧式VML図形) の中に文字を入れている。

検索語と期待される結果は EXPECTED / EXPECTED_ERRORS に定義し、tests/check_testdata.py で照合する。
旧形式 (.doc/.xls/.ppt) と PDF は LibreOffice (soffice) で新形式から変換して作る。
"""
from __future__ import annotations

import io
import os
import random
import shutil
import subprocess
import sys
import zipfile

import docx
import openpyxl
import pptx
from docx.oxml import parse_xml
from docx.shared import Pt
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, PatternFill, Side
from pptx.util import Inches

# ---------------------------------------------------------------------------
# 期待結果: (相対パス, 検索語, 種別, 場所, 行, 列, 位置ラベルに含まれる文字)
#   行・列が None のものは「－」表示 (図形内など) を期待する
#   列に "?" を置いたものは値を問わない
EXPECTED = [
    # ---- Word (docx)
    ("01_基本設計/基本設計書_受注管理.docx", "在庫引当処理", "word", "本文", 6, 5, "約1ページ"),
    ("01_基本設計/基本設計書_受注管理.docx", "ORD_STATUS", "word", "表1 (3行,2列)", 19, 1, "約2ページ"),
    ("01_基本設計/基本設計書_受注管理.docx", "ステータス定義", "word", "本文", 11, 4, "約2ページ"),
    ("01_基本設計/基本設計書_受注管理.docx", "図形内_承認フロー", "word", "図形", "?", "?", "図形: 承認フロー枠"),
    ("01_基本設計/基本設計書_受注管理.docx", "図形内_旧式テキスト", "word", "図形", "?", "?", "図形"),
    ("01_基本設計/基本設計書_受注管理.docx", "図形内_グループ注記", "word", "図形", "?", "?", "図形"),
    ("01_基本設計/基本設計書_受注管理.docx", "社外秘_受注管理", "word", "ヘッダー1", 1, 1, ""),
    ("01_基本設計/基本設計書_受注管理.docx", "コメント_要レビュー", "word", "コメント", 1, "?", ""),
    # ---- Word (doc 変換)
    ("01_基本設計/旧形式/基本設計書_受注管理.doc", "在庫引当処理", "word", "本文", "?", 5, ""),
    ("01_基本設計/旧形式/基本設計書_受注管理.doc", "ORD_STATUS", "word", "本文", "?", 1, ""),
    ("01_基本設計/旧形式/基本設計書_受注管理.doc", "図形内_承認フロー", "word", "図形", "?", "?", ""),
    ("01_基本設計/旧形式/基本設計書_受注管理.doc", "社外秘_受注管理", "word", "ヘッダー/フッター", "?", "?", ""),
    # ---- PDF (docx から変換)
    ("01_基本設計/旧形式/基本設計書_受注管理.pdf", "在庫引当処理", "pdf", "1ページ", "?", "?", ""),
    ("01_基本設計/旧形式/基本設計書_受注管理.pdf", "図形内_承認フロー", "pdf", "?", "?", "?", ""),
    # ---- Excel (xlsx)
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "顧客コード", "excel", "項目定義", 6, 3, "C6"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "図形内_画面遷移", "excel", "画面レイアウト", None, None, "図形: 遷移説明"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "図形内_グループ内", "excel", "画面レイアウト", None, None, "図形"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "図形内_吹き出し", "excel", "画面レイアウト", None, None, "図形: 吹き出し"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "非表示_旧仕様メモ", "excel", "旧仕様(非表示)", 2, 2, "B2"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "コメント_要確認", "excel", "項目定義", 6, 3, "コメント"),
    ("02_画面設計/受注 画面設計書 (v1.2).xlsx", "ORD-202610", "excel", "項目定義", 12, 5, "E12"),
    ("03_DB設計/テーブル定義書.xlsx", "ORD_STATUS", "excel", "T_ORDER", 9, 2, "B9"),
    # ---- Excel (xls 変換)
    ("02_画面設計/旧形式/受注 画面設計書 (v1.2).xls", "顧客コード", "excel", "項目定義", 6, 3, "C6"),
    ("02_画面設計/旧形式/受注 画面設計書 (v1.2).xls", "図形内_画面遷移", "excel", "画面レイアウト", None, None, "図形"),
    ("02_画面設計/旧形式/受注 画面設計書 (v1.2).xls", "非表示_旧仕様メモ", "excel", "旧仕様(非表示)", 2, 2, "B2"),
    ("03_DB設計/旧形式/テーブル定義書.xls", "ORD_STATUS", "excel", "T_ORDER", 9, 2, "B9"),
    # ---- PowerPoint
    ("04_説明資料/受注管理_方式説明.pptx", "性能要件_3秒以内", "ppt", "スライド2", "?", "?", ""),
    ("04_説明資料/受注管理_方式説明.pptx", "図形内_構成図ラベル", "ppt", "スライド3", "?", "?", ""),
    ("04_説明資料/受注管理_方式説明.pptx", "ノート_補足説明", "ppt", "スライド1 ノート", 1, 1, ""),
    ("04_説明資料/旧形式/受注管理_方式説明.ppt", "性能要件_3秒以内", "ppt", "スライド2", "?", "?", ""),
    ("04_説明資料/旧形式/受注管理_方式説明.ppt", "図形内_構成図ラベル", "ppt", "スライド3", "?", "?", ""),
    ("04_説明資料/旧形式/受注管理_方式説明.ppt", "ノート_補足説明", "ppt", "スライド1 ノート", 1, 1, ""),
    # ---- テキスト (ソース・設定・データ)
    ("10_src/main/java/jp/example/order/OrderService.java", "calcOrderAmount", "text", "", 23, 23, ""),
    ("10_src/main/resources/sql/create_t_order.sql", "ORD_STATUS", "text", "", 6, 5, ""),
    ("10_src/main/resources/config/application.xml", "order.timeout", "text", "", 5, 17, ""),
    ("10_src/main/resources/messages_ja.properties", "受注番号が不正です", "text", "", 3, 20, ""),
    ("10_src/main/webapp/js/order.js", "calcOrderAmount", "text", "", 4, 10, ""),
    ("20_data/受注データ_SJIS.csv", "ORD-20261012-0003", "text", "", 4, 1, ""),
    ("20_data/受注データ_UTF8BOM.csv", "ORD-20261012-0103", "text", "", 4, 1, ""),
    ("20_data/legacy_EUC.txt", "在庫引当処理", "text", "", 2, 7, ""),
    ("20_data/batch_UTF16.log", "在庫引当処理", "text", "", 3, 27, ""),
    ("20_data/long_line.txt", "LONGLINE_MARKER", "text", "", 1, 50001, ""),
    ("20_data/Makefile", "ORD_STATUS", "text", "", 2, 7, ""),
]

# 「読めなかったファイル」に出るべきもの: (相対パス, 理由に含まれる文字)
EXPECTED_ERRORS = [
    ("90_異常系/パスワード付き_受注一覧.xlsx", "パスワード"),
    ("90_異常系/パスワード付き_仕様書.docx", "パスワード"),
    ("90_異常系/破損ファイル.xlsx", "壊れて"),
    ("90_異常系/破損ファイル.docx", "壊れて"),
    ("90_異常系/スキャン画像のみ.pdf", "テキストが含まれていません"),
]
# 検索対象にならない (結果にも読めなかったファイルにも出ない) べきもの
EXPECTED_NOT_FOUND = [
    "90_異常系/~$受注 画面設計書.xlsx",   # Office の一時ファイル
    "90_異常系/order_dump.bin",          # バイナリ拡張子
    ".git/COMMIT_EDITMSG",              # 除外フォルダ
]

W_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office"')


def _txbx(lines):
    return "<w:txbxContent>" + "".join(
        f"<w:p><w:r><w:t xml:space=\"preserve\">{t}</w:t></w:r></w:p>" for t in lines) + "</w:txbxContent>"


def word_textbox_run(doc_pr_id, name, lines, y_emu=0):
    """Word 2010 以降が保存する形 (DrawingML 図形 + 旧式 VML の代替表現) のテキストボックス"""
    return parse_xml(f"""
<w:r {W_NS}><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>
 <wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="{251659264 + doc_pr_id}"
   behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/>
  <wp:positionH relativeFrom="column"><wp:posOffset>3200400</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="paragraph"><wp:posOffset>{y_emu}</wp:posOffset></wp:positionV>
  <wp:extent cx="2286000" cy="685800"/><wp:effectExtent l="0" t="0" r="0" b="0"/>
  <wp:wrapSquare wrapText="bothSides"/>
  <wp:docPr id="{doc_pr_id}" name="{name}"/><wp:cNvGraphicFramePr/>
  <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
   <wps:wsp><wps:cNvSpPr txBox="1"/>
    <wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2286000" cy="685800"/></a:xfrm>
     <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFF2CC"/></a:solidFill>
     <a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr>
    <wps:txbx>{_txbx(lines)}</wps:txbx>
    <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" anchor="t"/>
   </wps:wsp>
  </a:graphicData></a:graphic>
 </wp:anchor></w:drawing></mc:Choice>
 <mc:Fallback><w:pict>
  <v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe"/>
  <v:shape id="Text Box {doc_pr_id}" o:spid="_x0000_s10{doc_pr_id}" type="#_x0000_t202"
    style="position:absolute;margin-left:252pt;margin-top:0;width:180pt;height:54pt;z-index:{doc_pr_id}" fillcolor="#fff2cc">
   <v:textbox>{_txbx(lines)}</v:textbox>
  </v:shape></w:pict></mc:Fallback>
</mc:AlternateContent></w:r>""")


def word_vml_run(lines):
    """Word 2003 時代の VML だけで書かれたテキストボックス (古い文書を docx 化したもの等)"""
    return parse_xml(f"""
<w:r {W_NS}><w:pict>
 <v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe"/>
 <v:shape id="OldBox1" type="#_x0000_t202" style="position:absolute;margin-left:0;margin-top:6pt;width:200pt;height:40pt;z-index:9">
  <v:textbox>{_txbx(lines)}</v:textbox>
 </v:shape></w:pict></w:r>""")


def word_group_run(doc_pr_id):
    """グループ化した図形 (四角形2つ、それぞれに文字)"""
    def sp(i, x, text):
        return f"""<wps:wsp><wps:cNvPr id="{100 + i}" name="四角形 {i}"/><wps:cNvSpPr/>
 <wps:spPr><a:xfrm><a:off x="{x}" y="0"/><a:ext cx="1371600" cy="457200"/></a:xfrm>
  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>
 <wps:txbx>{_txbx([text])}</wps:txbx><wps:bodyPr/></wps:wsp>"""
    return parse_xml(f"""
<w:r {W_NS}><w:drawing>
 <wp:inline distT="0" distB="0" distL="0" distR="0">
  <wp:extent cx="2971800" cy="457200"/><wp:effectExtent l="0" t="0" r="0" b="0"/>
  <wp:docPr id="{doc_pr_id}" name="グループ化 {doc_pr_id}"/><wp:cNvGraphicFramePr/>
  <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup">
   <wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2971800" cy="457200"/>
    <a:chOff x="0" y="0"/><a:chExt cx="2971800" cy="457200"/></a:xfrm></wpg:grpSpPr>
    {sp(1, 0, "受注登録")}{sp(2, 1600200, "図形内_グループ注記")}
   </wpg:wgp>
  </a:graphicData></a:graphic>
 </wp:inline></w:drawing></w:r>""")


# ---------------------------------------------------------------------------
def make_word(path):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "MS Mincho"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "社外秘_受注管理システム 基本設計書"
    sec.footer.paragraphs[0].text = "Copyright Example Corp."
    d.add_heading("基本設計書（受注管理システム）", 0)                       # 段落1
    d.add_paragraph("文書番号: DS-ORD-001　版数: 1.2　作成日: 2026/10/01")  # 段落2
    d.add_heading("1. 概要", 1)                                             # 段落3
    d.add_paragraph("本書は受注管理システムの基本設計を記述する。")         # 段落4
    d.add_heading("2. 処理概要", 1)                                         # 段落5
    p = d.add_paragraph("受注後に在庫引当処理を行い、引当結果を受注ステータスへ反映する。")  # 段落6
    d.add_comment(p.runs[0], text="コメント_要レビュー: 引当失敗時の扱いを追記すること", author="レビュー担当")
    p7 = d.add_paragraph("承認フローは右図のとおり。")                      # 段落7 (図形のアンカー)
    p7._p.append(word_textbox_run(11, "承認フロー枠", ["図形内_承認フロー", "申請 → 課長承認 → 部長承認"]))
    p8 = d.add_paragraph("旧版から引き継いだ注記:")                         # 段落8
    p8._p.append(word_vml_run(["図形内_旧式テキスト（Word2003 形式）"]))
    p9 = d.add_paragraph("処理の流れ: ")                                     # 段落9
    p9._p.append(word_group_run(12))
    d.add_page_break()                                                      # 段落10 (改ページ)
    d.add_heading("3. ステータス定義", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    rows = [("区分", "項目名", "説明"), ("01", "受注番号", "ORD-yyyymmdd-nnnn"),
            ("02", "ORD_STATUS", "受注ステータス (10:受付 20:引当済 30:出荷済)"), ("03", "顧客コード", "8桁")]
    for r, vals in enumerate(rows):
        for c, v in enumerate(vals):
            t.cell(r, c).text = v
    d.add_paragraph("以上")
    d.save(path)
    # Word が保存時に書く「表示上の改ページ位置」を真似る (ページ目安の確認用)
    _add_rendered_page_break(path)


def _add_rendered_page_break(path):
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    xml = items["word/document.xml"].decode("utf-8")
    key = "3. ステータス定義"
    i = xml.find(key)
    j = xml.rfind("<w:r>", 0, i)
    if j < 0:
        j = xml.rfind("<w:r ", 0, i)
    xml = xml[:j] + "<w:r><w:lastRenderedPageBreak/></w:r>" + xml[j:]
    items["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)


# ---------------------------------------------------------------------------
XDR = ('xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def _xsp(i, name, text_lines, prst="rect", fill="DDEBF7"):
    paras = "".join(f"<a:p><a:r><a:rPr lang=\"ja-JP\" sz=\"1100\"/><a:t>{t}</a:t></a:r></a:p>" for t in text_lines)
    return f"""<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{i}" name="{name}"/><xdr:cNvSpPr/></xdr:nvSpPr>
<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="457200"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>
<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="2F5597"/></a:solidFill></a:ln></xdr:spPr>
<xdr:txBody><a:bodyPr vertOverflow="clip" wrap="square" rtlCol="0" anchor="ctr"/><a:lstStyle/>{paras}</xdr:txBody></xdr:sp>"""


def _anchor(c1, r1, c2, r2, body):
    return f"""<xdr:twoCellAnchor><xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>{c2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>{body}<xdr:clientData/></xdr:twoCellAnchor>"""


def _add_excel_shapes(path, sheet_index):
    """openpyxl は図形を書けないので、保存後の xlsx に DrawingML の図形を追加する"""
    group = f"""<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="20" name="グループ 20"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>
<xdr:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="3657600" cy="457200"/><a:chOff x="0" y="0"/><a:chExt cx="3657600" cy="457200"/></a:xfrm></xdr:grpSpPr>
{_xsp(21, "四角形 21", ["受注一覧へ"])}{_xsp(22, "四角形 22", ["図形内_グループ内ボタン説明"])}</xdr:grpSp>"""
    drawing = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr {XDR}>
{_anchor(7, 2, 11, 5, _xsp(2, "遷移説明", ["図形内_画面遷移：登録ボタン押下で確認画面へ", "キャンセルで一覧へ戻る"]))}
{_anchor(7, 8, 13, 10, group)}
{_anchor(1, 14, 5, 17, _xsp(3, "吹き出し 3", ["図形内_吹き出し：必須項目は赤枠"], prst="wedgeRectCallout", fill="FFF2CC"))}
</xdr:wsDr>"""
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    sheet = f"xl/worksheets/sheet{sheet_index}.xml"
    rels = f"xl/worksheets/_rels/sheet{sheet_index}.xml.rels"
    items["xl/drawings/drawing1.xml"] = drawing.encode("utf-8")
    rel_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               '<Relationship Id="rIdDr1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
               'Target="../drawings/drawing1.xml"/></Relationships>')
    if rels in items:
        s = items[rels].decode("utf-8")
        s = s.replace("</Relationships>", rel_xml.split("<Relationships")[1].split(">", 1)[1].replace("</Relationships>", "") + "</Relationships>")
        items[rels] = s.encode("utf-8")
    else:
        items[rels] = rel_xml.encode("utf-8")
    s = items[sheet].decode("utf-8")
    if 'xmlns:r=' not in s.split(">", 2)[1]:
        s = s.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    # <drawing> は pageMargins/pageSetup の後、legacyDrawing の前に置く
    if "<legacyDrawing" in s:
        s = s.replace("<legacyDrawing", '<drawing r:id="rIdDr1"/><legacyDrawing', 1)
    else:
        s = s.replace("</worksheet>", '<drawing r:id="rIdDr1"/></worksheet>')
    items[sheet] = s.encode("utf-8")
    ct = items["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace("</Types>", '<Override PartName="/xl/drawings/drawing1.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    items["[Content_Types].xml"] = ct.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)


def make_screen_design(path):
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="D9E1F2")

    ws0 = wb.active
    ws0.title = "表紙"
    ws0["B2"] = "受注管理システム"
    ws0["B4"] = "画面設計書（受注登録画面）"
    ws0["B8"] = "版数"
    ws0["C8"] = "1.2"
    ws0.merge_cells("B4:H5")

    ws = wb.create_sheet("画面レイアウト")   # sheet2 → 図形を載せる
    for col in range(1, 40):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 2.7  # Excel方眼紙
    ws["B2"] = "受注登録画面"
    ws.merge_cells("B2:P2")
    ws["B4"] = "受注番号"
    ws["F4"] = "（自動採番）"
    ws["B6"] = "顧客"
    ws["F6"] = "[顧客コード] [検索]"
    ws["B12"] = "[登録] [キャンセル]"

    ws2 = wb.create_sheet("項目定義")
    hdr = ["No", "項目名(論理)", "項目名(物理)", "型", "桁", "必須", "備考"]
    for c, v in enumerate(hdr, 2):
        cell = ws2.cell(row=5, column=c, value=v)
        cell.fill, cell.border = head, border
    data = [
        (1, "顧客コード", "CUST_CD", "CHAR", 8, "○", "顧客マスタ参照"),
        (2, "受注日", "ORDER_DATE", "DATE", None, "○", "システム日付を初期表示"),
        (3, "受注ステータス", "ORD_STATUS", "CHAR", 2, "○", "10:受付"),
        (4, "数量", "QTY", "NUMBER", 5, "○", "1以上"),
        (5, "単価", "UNIT_PRICE", "NUMBER", 9, "", "税抜"),
        (6, "受注番号", "ORDER_NO", "CHAR", 17, "", "例: ORD-20261012-0001"),
    ]
    # B6 から: B=No C=論理名 D=物理名 ...  → 顧客コードは C6
    for r, row in enumerate(data, 6):
        for c, v in enumerate(row, 2):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2["C6"].comment = Comment("コメント_要確認: 桁数は10桁に変更予定", "設計者")
    ws2["B12"] = "採番例"
    ws2["E12"] = "ORD-20261012-0001"
    ws2["B14"] = 20261012          # 数値
    ws2["C14"] = 12345.5

    ws3 = wb.create_sheet("旧仕様")
    ws3["B2"] = "非表示_旧仕様メモ：顧客コードは6桁だった"
    ws3.sheet_state = "hidden"

    ws4 = wb.create_sheet("変更履歴")
    ws4.append(["版数", "日付", "内容", "担当"])
    ws4.append(["1.0", "2026/09/01", "新規作成", "山田"])
    ws4.append(["1.2", "2026/10/01", "顧客コードの桁数を変更", "佐藤"])
    wb.save(path)
    _add_excel_shapes(path, 2)


def make_table_def(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T_ORDER"
    ws["B2"] = "テーブル定義書"
    ws["B3"] = "テーブル名"
    ws["C3"] = "T_ORDER（受注）"
    ws.append([])
    cols = [("ORDER_NO", "受注番号", "CHAR(17)", "PK"), ("CUST_CD", "顧客コード", "CHAR(8)", ""),
            ("ORDER_DATE", "受注日", "DATE", ""), ("ORD_STATUS", "受注ステータス", "CHAR(2)", ""),
            ("QTY", "数量", "NUMBER(5)", ""), ("UNIT_PRICE", "単価", "NUMBER(9)", "")]
    ws["B5"], ws["C5"], ws["D5"], ws["E5"] = "物理名", "論理名", "型", "キー"
    for r, (a, b, c, d) in enumerate(cols, 6):
        ws.cell(row=r, column=2, value=a)
        ws.cell(row=r, column=3, value=b)
        ws.cell(row=r, column=4, value=c)
        ws.cell(row=r, column=5, value=d)
    # ORD_STATUS は B9
    ws2 = wb.create_sheet("M_CUSTOMER")
    ws2.append(["物理名", "論理名"])
    for i in range(1, 2001):  # 行数の多いシート
        ws2.append([f"COL_{i:04d}", f"項目{i}"])
    wb.save(path)


def make_ppt(path):
    p = pptx.Presentation()
    s1 = p.slides.add_slide(p.slide_layouts[0])
    s1.shapes.title.text = "受注管理システム 方式説明"
    s1.placeholders[1].text = "2026年10月　開発チーム"
    s1.notes_slide.notes_text_frame.text = "ノート_補足説明：冒頭で背景を説明する"
    s2 = p.slides.add_slide(p.slide_layouts[5])
    s2.shapes.title.text = "非機能要件"
    tbl = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(7), Inches(1.5)).table
    for r, (a, b) in enumerate([("項目", "要件"), ("性能", "性能要件_3秒以内（画面応答）"), ("可用性", "99.5%")]):
        tbl.cell(r, 0).text = a
        tbl.cell(r, 1).text = b
    s3 = p.slides.add_slide(p.slide_layouts[5])
    s3.shapes.title.text = "システム構成"
    grp = s3.shapes.add_group_shape()
    b1 = grp.shapes.add_shape(1, Inches(1), Inches(2), Inches(2), Inches(1))
    b1.text_frame.text = "Webサーバ"
    b2 = grp.shapes.add_shape(1, Inches(4), Inches(2), Inches(3), Inches(1))
    b2.text_frame.text = "図形内_構成図ラベル（APサーバ）"
    p.save(path)


# ---------------------------------------------------------------------------
def make_texts(out):
    def w(rel, text, enc="utf-8", newline="\n"):
        path = os.path.join(out, rel)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding=enc, newline=newline) as f:
            f.write(text)

    java = ["package jp.example.order;", "", "import java.math.BigDecimal;", "import java.util.List;", "",
            "/**", " * 受注サービス", " * 在庫引当処理を呼び出す。", " */",
            "public class OrderService {", "", "    private final StockService stockService;", "",
            "    public OrderService(StockService stockService) {",
            "        this.stockService = stockService;", "    }", "",
            "    /**", "     * 受注金額を計算する。", "     * @param order 受注", "     * @return 金額", "     */",
            "    public BigDecimal calcOrderAmount(Order order) {",           # 23行目, 23桁目
            "        return order.getUnitPrice().multiply(BigDecimal.valueOf(order.getQty()));",
            "    }", "}", ""]
    w("10_src/main/java/jp/example/order/OrderService.java", "\r\n".join(java), newline="")
    w("10_src/main/resources/sql/create_t_order.sql",
      "-- 受注テーブル\nCREATE TABLE T_ORDER (\n    ORDER_NO   CHAR(17) NOT NULL,\n    CUST_CD    CHAR(8)  NOT NULL,\n"
      "    ORDER_DATE DATE,\n    ORD_STATUS CHAR(2)  DEFAULT '10',\n    CONSTRAINT PK_T_ORDER PRIMARY KEY (ORDER_NO)\n);\n")
    w("10_src/main/resources/config/application.xml",
      '<?xml version="1.0" encoding="UTF-8"?>\n<config>\n  <!-- 受注設定 -->\n  <order>\n'
      '    <entry key="order.timeout">3000</entry>\n  </order>\n</config>\n')
    w("10_src/main/resources/messages_ja.properties",
      "# メッセージ\nerror.required={0}は必須です\nerror.orderNo=E001:受注番号が不正です\n", enc="utf-8")
    w("10_src/main/webapp/js/order.js",
      "'use strict';\n\n// 受注画面\nfunction calcOrderAmount(qty, price) {\n  return qty * price;\n}\n")
    w("10_src/main/webapp/order.html",
      "<!DOCTYPE html>\n<html lang=\"ja\"><head><meta charset=\"utf-8\"><title>受注登録</title></head>\n"
      "<body><h1>受注登録</h1><label>顧客コード</label><input id=\"custCd\"></body></html>\n")
    csv_rows = ["受注番号,顧客コード,受注日,ORD_STATUS,数量"] + \
               [f"ORD-20261012-{i:04d},C{1000 + i:07d},2026/10/12,10,{i * 3}" for i in range(1, 6)]
    w("20_data/受注データ_SJIS.csv", "\r\n".join(csv_rows) + "\r\n", enc="cp932", newline="")
    csv2 = ["受注番号,顧客名"] + [f"ORD-20261012-{100 + i:04d},株式会社テスト{i}" for i in range(1, 6)]
    w("20_data/受注データ_UTF8BOM.csv", "\r\n".join(csv2) + "\r\n", enc="utf-8-sig", newline="")
    w("20_data/legacy_EUC.txt", "旧システムの仕様メモ\n夜間バッチで在庫引当処理を実施する。\n", enc="euc_jp")
    w("20_data/batch_UTF16.log",
      "2026-10-12 01:00:00 INFO  バッチ開始\n2026-10-12 01:00:01 INFO  受注取込 120件\n"
      "2026-10-12 01:00:05 INFO  在庫引当処理 完了\n", enc="utf-16")
    w("20_data/long_line.txt", "x" * 50000 + "LONGLINE_MARKER" + "y" * 50000 + "\n")
    w("20_data/Makefile", "build:\n\techo ORD_STATUS check\n")
    w("20_data/空ファイル.txt", "")
    w(".git/COMMIT_EDITMSG", "fix: ORD_STATUS の初期値\n")
    # 大量のソース (性能確認用・ヒットしない語ばかり)
    random.seed(7)
    words = ["order", "stock", "customer", "amount", "status", "update", "select", "insert", "commit", "受付"]
    for i in range(300):
        body = "\n".join(" ".join(random.choice(words) for _ in range(10)) for _ in range(200))
        w(f"10_src/main/java/jp/example/gen/Gen{i:03d}.java", f"// generated {i}\n{body}\n")


def make_abnormal(out, screen_xlsx, word_docx):
    d = os.path.join(out, "90_異常系")
    os.makedirs(d, exist_ok=True)
    # パスワード付き (OOXML 暗号化)
    try:
        from msoffcrypto.format.ooxml import OOXMLFile
        for src, name in ((screen_xlsx, "パスワード付き_受注一覧.xlsx"), (word_docx, "パスワード付き_仕様書.docx")):
            with open(src, "rb") as f:
                plain = f.read()
            enc_out = io.BytesIO()
            OOXMLFile(io.BytesIO(plain)).encrypt("pass1234", enc_out)
            with open(os.path.join(d, name), "wb") as f:
                f.write(enc_out.getvalue())
    except Exception as e:  # msoffcrypto-tool が無い場合
        print("パスワード付きファイルは作成できませんでした:", e)
    # 破損 (途中で切れた ZIP)
    for src, name in ((screen_xlsx, "破損ファイル.xlsx"), (word_docx, "破損ファイル.docx")):
        with open(src, "rb") as f:
            data = f.read()
        with open(os.path.join(d, name), "wb") as f:
            f.write(data[: len(data) // 3])
    # 画像だけの PDF (文字情報なし)
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(os.path.join(d, "スキャン画像のみ.pdf"), pagesize=A4)
    for i in range(30):
        c.rect(50 + i * 5, 700 - i * 10, 200, 8, fill=1)
    c.showPage()
    c.save()
    # Office の一時ファイル / バイナリ
    with open(os.path.join(d, "~$受注 画面設計書.xlsx"), "wb") as f:
        f.write(b"\x0cuser" + b"\x00" * 150)
    with open(os.path.join(d, "order_dump.bin"), "wb") as f:
        f.write(b"\x00\x01ORD_STATUS\x00" * 100)


def soffice(src, fmt, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "testdata"
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    j = lambda *p: os.path.join(out, *p)
    for d in ("01_基本設計", "02_画面設計", "03_DB設計", "04_説明資料"):
        os.makedirs(j(d), exist_ok=True)
    word = j("01_基本設計", "基本設計書_受注管理.docx")
    screen = j("02_画面設計", "受注 画面設計書 (v1.2).xlsx")
    tdef = j("03_DB設計", "テーブル定義書.xlsx")
    ppt = j("04_説明資料", "受注管理_方式説明.pptx")
    make_word(word)
    make_screen_design(screen)
    make_table_def(tdef)
    make_ppt(ppt)
    make_texts(out)
    make_abnormal(out, screen, word)
    if shutil.which("soffice"):
        soffice(word, "doc", j("01_基本設計", "旧形式"))
        soffice(word, "pdf", j("01_基本設計", "旧形式"))
        soffice(screen, "xls", j("02_画面設計", "旧形式"))
        soffice(tdef, "xls", j("03_DB設計", "旧形式"))
        soffice(ppt, "ppt", j("04_説明資料", "旧形式"))
    else:
        print("soffice が無いため旧形式・PDF は作成しませんでした")
    n = sum(len(f) for _, _, f in os.walk(out))
    print(f"作成しました: {out} ({n} ファイル)")


if __name__ == "__main__":
    main()
