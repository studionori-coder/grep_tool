"""検索結果を Excel (.xlsx) に出力する。ファイル名・フォルダはクリックで開けるリンクにする。"""
from __future__ import annotations

import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .extractors import ENCODING_DISPLAY, KIND_LABEL

try:  # openpyxl 3.1 以降: セル内の一部だけ色を付けられる
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    _RICH = True
except Exception:  # pragma: no cover
    _RICH = False

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_CELL = 32000


def _clean(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub(" ", v)
        if len(v) > MAX_CELL:
            v = v[:MAX_CELL] + "…"
    return v


def _file_url(path: str) -> str:
    p = os.path.abspath(path)
    if p.startswith("\\\\"):
        return "file:///" + p
    if os.sep == "\\":
        return "file:///" + p
    return "file://" + p


def _set_text(cell, value):
    cell.value = _clean(value)
    if isinstance(value, str):
        cell.data_type = "s"  # "=" で始まる文字列を数式と解釈させない


def _link(cell, path: str, text: str, location: str | None = None):
    _set_text(cell, text)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=_file_url(path), location=location,
                               tooltip="クリックで開く")
    cell.font = LINK_FONT


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _fmt_time(ts: float) -> str:
    return dt.datetime.fromtimestamp(ts).strftime("%Y/%m/%d %H:%M") if ts else ""


def _excel_location(place: str, label: str) -> str | None:
    import re
    m = re.match(r"^([A-Z]+\d+)", label or "")
    if not m:
        return None
    sheet = re.sub(r"\(非表示\)$", "", place).replace("'", "''")
    return f"'{sheet}'!{m.group(1)}"


def export(path: str, hits, files, conditions: list[tuple[str, str]], summary_lines: list[str],
           progress=None):
    """hits: 検索結果 (表示順), files: FileInfo のリスト (index で参照)"""
    wb = Workbook()

    # ---- 検索結果
    ws = wb.active
    ws.title = "検索結果"
    headers = ["No", "ファイル名", "種別", "場所", "行", "列", "位置(補足)", "ヒット文字列",
               "該当箇所", "件数", "フォルダ", "フルパス", "更新日時", "文字コード"]
    _header(ws, headers, [7, 30, 11, 18, 7, 7, 14, 16, 70, 6, 40, 60, 16, 10])
    hl_font = InlineFont(b=True, color="C00000") if _RICH else None
    n = len(hits)
    for i, h in enumerate(hits, 1):
        fidx, kind, place, line, col, label, match, ctx, hs, he, cnt = h
        f = files[fidx]
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        loc = _excel_location(place, label) if kind == "excel" and f.path.lower().endswith(
            (".xlsx", ".xlsm", ".xls")) else None
        _link(ws.cell(row=r, column=2), f.path, os.path.basename(f.path), loc)
        _set_text(ws.cell(row=r, column=3), KIND_LABEL.get(kind, kind))
        _set_text(ws.cell(row=r, column=4), place)
        ws.cell(row=r, column=5, value=line if line is not None else "－")
        ws.cell(row=r, column=6, value=col if col is not None else "－")
        _set_text(ws.cell(row=r, column=7), label)
        _set_text(ws.cell(row=r, column=8), match)
        c = ws.cell(row=r, column=9)
        ctx = _clean(ctx)
        if _RICH and 0 <= hs < he <= len(ctx):
            parts = []
            if hs > 0:
                parts.append(ctx[:hs])
            parts.append(TextBlock(hl_font, ctx[hs:he]))
            if he < len(ctx):
                parts.append(ctx[he:])
            c.value = CellRichText(parts)
        else:
            _set_text(c, ctx)
        ws.cell(row=r, column=10, value=cnt)
        _link(ws.cell(row=r, column=11), os.path.dirname(f.path), os.path.dirname(f.path))
        _set_text(ws.cell(row=r, column=12), f.path)
        _set_text(ws.cell(row=r, column=13), _fmt_time(f.mtime))
        _set_text(ws.cell(row=r, column=14), ENCODING_DISPLAY.get(f.encoding, f.encoding))
        if progress and i % 2000 == 0:
            progress(i, n)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, n + 1)}"

    # ---- ファイル別
    ws2 = wb.create_sheet("ファイル別")
    _header(ws2, ["ファイル名", "種別", "ヒット件数", "フォルダ", "フルパス", "サイズ(KB)", "更新日時", "文字コード", "備考"],
            [32, 11, 10, 40, 60, 11, 16, 10, 30])
    r = 1
    for f in files:
        if f.hits <= 0:
            continue
        r += 1
        _link(ws2.cell(row=r, column=1), f.path, os.path.basename(f.path))
        _set_text(ws2.cell(row=r, column=2), KIND_LABEL.get(f.kind, f.kind))
        ws2.cell(row=r, column=3, value=f.hits)
        _link(ws2.cell(row=r, column=4), os.path.dirname(f.path), os.path.dirname(f.path))
        _set_text(ws2.cell(row=r, column=5), f.path)
        ws2.cell(row=r, column=6, value=round(f.size / 1024, 1))
        _set_text(ws2.cell(row=r, column=7), _fmt_time(f.mtime))
        _set_text(ws2.cell(row=r, column=8), ENCODING_DISPLAY.get(f.encoding, f.encoding))
        _set_text(ws2.cell(row=r, column=9), "ヒットが多いため途中で打ち切り" if f.truncated else "")
    ws2.auto_filter.ref = f"A1:I{max(1, r)}"

    # ---- 読めなかったファイル
    ws3 = wb.create_sheet("読めなかったファイル")
    _header(ws3, ["ファイル名", "種別", "状態", "理由", "フルパス"], [32, 11, 10, 50, 60])
    r = 1
    for f in files:
        if f.status in ("error", "skip"):
            r += 1
            _link(ws3.cell(row=r, column=1), f.path, os.path.basename(f.path))
            _set_text(ws3.cell(row=r, column=2), KIND_LABEL.get(f.kind, f.kind))
            _set_text(ws3.cell(row=r, column=3), "エラー" if f.status == "error" else "スキップ")
            _set_text(ws3.cell(row=r, column=4), f.message)
            _set_text(ws3.cell(row=r, column=5), f.path)

    # ---- 検索条件
    ws4 = wb.create_sheet("検索条件")
    _header(ws4, ["項目", "内容"], [22, 90])
    for k, v in conditions:
        ws4.append([_clean(k), _clean(v)])
        ws4.cell(row=ws4.max_row, column=2).data_type = "s"
    ws4.append([])
    for s in summary_lines:
        ws4.append(["", _clean(s)])
    ws4.append([])
    ws4.append(["行・列の見方", "テキスト: 行番号 / 行頭からの文字位置(桁)"])
    ws4.append(["", "Excel: セルの行番号 / 列番号 (位置欄にセル番地)"])
    ws4.append(["", "Word: 本文の段落番号 / 段落内の文字位置 (ページは目安)"])
    ws4.append(["", "PowerPoint: スライド内の行番号 / 文字位置"])
    ws4.append(["", "PDF: ページ内の行番号 / 文字位置"])
    ws4.append(["", "「－」は図形・テキストボックス等で行・列が無い箇所"])

    wb.save(path)
